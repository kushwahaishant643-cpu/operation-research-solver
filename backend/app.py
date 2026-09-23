from flask import Flask, jsonify, request, send_file
from flask_cors import CORS

from algorithms.lpp import solve_graphical_lpp
from algorithms.transportation import (
    north_west_corner,
    least_cost_method,
    vogel_approximation_method,
    modi_method
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import io
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image as XLImage


app = Flask(__name__)
CORS(app)


@app.route("/")
def home():
    return jsonify({
        "message": "Operation Research Smart Solver API is running!",
        "status": "success"
    })


@app.route("/api/test")
def test_api():
    return jsonify({
        "message": "Backend connection successful!",
        "project": "Operation Research Smart Solver"
    })


# ============================================================
# LPP - GRAPHICAL METHOD
# ============================================================

@app.route("/api/lpp/solve", methods=["POST"])
def solve_lpp():
    try:
        data = request.get_json()

        if not data:
            return jsonify({
                "status": "error",
                "message": "No JSON data was provided."
            }), 400

        objective = data.get("objective")
        constraints = data.get("constraints")
        optimization = data.get("optimization", "max")

        if objective is None:
            return jsonify({
                "status": "error",
                "message": "Objective function is required."
            }), 400

        if constraints is None:
            return jsonify({
                "status": "error",
                "message": "Constraints are required."
            }), 400

        result = solve_graphical_lpp(
            objective_coefficients=objective,
            constraints=constraints,
            optimization=optimization
        )

        return jsonify(result), 200

    except ValueError as error:
        return jsonify({
            "status": "error",
            "message": str(error)
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 500



# ============================================================
# LPP - WORD PROBLEM PARSER
# ============================================================

def _num(v):
    try:
        return float(str(v).replace(",", "").replace("₹", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def _parse_word_problem(statement):
    """Parse common two-variable LPP word problems into a mathematical model."""
    if not statement or not str(statement).strip():
        raise ValueError("Please enter a word problem statement.")

    text = " ".join(str(statement).replace("\n", " ").split())
    lower = text.lower()
    number = r"[0-9]+(?:\.[0-9]+)?"

    def num(value):
        return _num(value)

    def clean_name(value):
        value = re.sub(r"[^A-Za-z0-9_-]+$", "", str(value)).strip()
        return value.title()

    def norm_resource(value):
        value = re.sub(r"\b(hours?|hrs?|minutes?|mins?|days?|units?)\b", "", value, flags=re.I)
        value = re.sub(r"\s+", " ", value).strip(" ,.")
        return clean_name(value)

    # ------------------------------------------------------------
    # 1. Optimization direction
    # ------------------------------------------------------------
    optimization = "min" if re.search(
        r"\b(minimize|minimise|minimizing|minimising|minimum|least\s+cost)\b",
        lower,
    ) else "max"

    # ------------------------------------------------------------
    # 2. Detect objective coefficients.
    # Supports normal exam wording such as:
    # "profit earned on each unit of Product A is 40"
    # "profit per Product A is 40"
    # "Product A gives a profit of 40"
    # ------------------------------------------------------------
    objective_pairs = []
    objective_patterns = [
        rf"(?:profit|revenue|contribution|income|cost)\s+(?:earned|made|generated)?\s*(?:on|from|for|per)\s+(?:each\s+)?(?:unit\s+of\s+)?([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:is|=|of)\s*(?:₹|rs\.?|\$)?\s*({number})",
        rf"(?:profit|revenue|contribution|income|cost)\s+(?:per|for)\s+(?:each\s+)?(?:unit\s+of\s+)?([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:is|=|of)\s*(?:₹|rs\.?|\$)?\s*({number})",
        rf"(?:profit|revenue|contribution|income)\s+(?:of|from)\s+([A-Za-z][\w-]*)\s+(?:is|=)\s*(?:₹|rs\.?|\$)?\s*({number})",
        rf"(?:each|one)\s+([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:gives|earns|yields|generates)\s+(?:a\s+)?(?:profit|revenue|contribution|income)\s+(?:of|is|=)?\s*(?:₹|rs\.?|\$)?\s*({number})",
        rf"([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:gives|earns|yields|generates)\s+(?:a\s+)?(?:profit|revenue|contribution|income)\s+(?:of|is|=)?\s*(?:₹|rs\.?|\$)?\s*({number})",
    ]

    for pattern in objective_patterns:
        for match in re.finditer(pattern, text, re.I):
            product = clean_name(match.group(1))
            value = num(match.group(2))
            if product and value is not None and product.lower() not in {p.lower() for p, _ in objective_pairs}:
                objective_pairs.append((product, value))

    # ------------------------------------------------------------
    # 3. Detect product requirements.
    # Example:
    # "Each Product A requires 2 hours of machine time and 1 hour of labour"
    # ------------------------------------------------------------
    product_requirements = {}
    req_pattern = re.compile(
        rf"(?:each\s+|one\s+|a\s+)?([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:requires|needs|uses|consumes|takes)\s+"
        rf"(.+?)(?=\s+(?:while|whereas|and\s+(?:each|one|a)\s+[A-Za-z][\w-]*\s+(?:requires|needs|uses|consumes|takes))\b|[.;]|$)",
        re.I,
    )

    for match in req_pattern.finditer(text):
        product = clean_name(match.group(1))
        body = match.group(2).strip(" ,.")
        requirements = []
        req_value_pattern = re.compile(
            rf"({number})\s+(?:hours?|hrs?|minutes?|mins?|kgs?|kg|kilograms?|litres?|liters?|tons?|tonnes?|units?)?\s*(?:of\s+)?([A-Za-z][A-Za-z0-9_-]*(?:\s+[A-Za-z][A-Za-z0-9_-]*){{0,3}}?)(?=\s+and\s+{number}\b|\s*,\s*{number}\b|$)",
            re.I,
        )
        for rm in req_value_pattern.finditer(body):
            value = num(rm.group(1))
            resource = norm_resource(rm.group(2))
            if value is not None and resource:
                if resource.lower() not in {r.lower() for r, _ in requirements}:
                    requirements.append((resource, value))
        if requirements:
            product_requirements[product] = requirements

    # Fallback for requirement sentences where the look-ahead is broader.
    if len(product_requirements) < 2:
        broad = re.finditer(
            rf"(?:each\s+|one\s+|a\s+)?([A-Za-z][\w-]*(?:\s+[A-Za-z][\w-]*)?)\s+(?:requires|needs|uses|consumes|takes)\s+([^.;]+)",
            text,
            re.I,
        )
        for match in broad:
            product = clean_name(match.group(1))
            body = match.group(2)
            requirements = []
            for rm in re.finditer(
                rf"({number})\s+(?:hours?|hrs?|minutes?|mins?|kgs?|kg|kilograms?|litres?|liters?|tons?|tonnes?|units?)?\s*(?:of\s+)?([A-Za-z][A-Za-z0-9_-]*(?:\s+[A-Za-z][A-Za-z0-9_-]*){{0,3}}?)(?=\s+and\s+{number}\b|$)",
                body,
                re.I,
            ):
                value = num(rm.group(1))
                resource = norm_resource(rm.group(2))
                if value is not None and resource:
                    requirements.append((resource, value))
            if requirements:
                product_requirements[product] = requirements

    # ------------------------------------------------------------
    # 4. Establish product order.
    # ------------------------------------------------------------
    products = []
    for product, _ in objective_pairs:
        if product.lower() not in {p.lower() for p in products}:
            products.append(product)
    for product in product_requirements:
        if product.lower() not in {p.lower() for p in products}:
            products.append(product)
    products = products[:2]

    # ------------------------------------------------------------
    # 5. Objective coefficients in product order.
    # ------------------------------------------------------------
    objective_map = {p.lower(): c for p, c in objective_pairs}
    objective_coefficients = [objective_map.get(p.lower()) for p in products]
    if len(objective_coefficients) < 2:
        objective_coefficients += [None] * (2 - len(objective_coefficients))

    # ------------------------------------------------------------
    # 6. Build resource matrix.
    # ------------------------------------------------------------
    resource_names = []
    for product in products:
        for resource, _ in product_requirements.get(product, []):
            if resource.lower() not in {r.lower() for r in resource_names}:
                resource_names.append(resource)

    resources = []
    for resource in resource_names:
        coeffs = []
        for product in products:
            value = None
            for rname, rvalue in product_requirements.get(product, []):
                if rname.lower() == resource.lower() or rname.lower() in resource.lower() or resource.lower() in rname.lower():
                    value = rvalue
                    break
            coeffs.append(value if value is not None else 0)
        resources.append({"resource": resource, "coefficients": coeffs, "rhs": None, "operator": "<="})

    # ------------------------------------------------------------
    # 7. Detect resource availability.
    # Examples:
    # "100 machine hours are available"
    # "80 labour hours are available"
    # "maximum of 100 machine hours"
    # ------------------------------------------------------------
    availability_matches = list(re.finditer(
        rf"({number})\s+(?:hours?|hrs?|kgs?|kg|kilograms?|litres?|liters?|tons?|tonnes?|units?)?\s*(?:of\s+)?([A-Za-z][A-Za-z0-9_-]*(?:\s+[A-Za-z][A-Za-z0-9_-]*){{0,3}}?)\s+(?:are|is)\s+(?:available|available\s+per\s+week|provided|present)",
        text,
        re.I,
    ))

    # Also catch "100 machine hours are available" where resource comes before unit.
    availability_matches += list(re.finditer(
        rf"({number})\s+([A-Za-z][A-Za-z0-9_-]*)\s+(?:hours?|hrs?|kgs?|kg|kilograms?|litres?|liters?|tons?|tonnes?|units?)\s+(?:are|is)\s+(?:available|provided|present)",
        text,
        re.I,
    ))

    for match in availability_matches:
        value = num(match.group(1))
        raw_resource = norm_resource(match.group(2))
        if value is None or not raw_resource:
            continue
        best = None
        best_score = 0
        for resource in resources:
            rt = set(resource["resource"].lower().split())
            at = set(raw_resource.lower().split())
            score = len(rt & at)
            if resource["resource"].lower() == raw_resource.lower():
                score += 10
            if score > best_score:
                best_score = score
                best = resource
        if best is not None and best_score > 0:
            best["rhs"] = value
            best["operator"] = "<="

    # Explicit resource-first forms.
    for resource in resources:
        r = re.escape(resource["resource"])
        patterns = [
            (rf"{r}\s+(?:hours?|hrs?)?\s*(?:are|is)?\s*(?:available|availability)\s*(?:is|=|of|:)?\s*({number})", "<="),
            (rf"({number})\s+(?:hours?|hrs?)?\s+(?:of\s+)?{r}\s+(?:are|is)\s+available", "<="),
            (rf"(?:at\s+most|no\s+more\s+than|maximum\s+of)\s+({number})\s+(?:hours?|hrs?)?\s+(?:of\s+)?{r}", "<="),
        ]
        for pattern, operator in patterns:
            if resource["rhs"] is None:
                m = re.search(pattern, text, re.I)
                if m:
                    resource["rhs"] = num(m.group(1))
                    resource["operator"] = operator
                    break

    ready = (
        len(products) == 2
        and all(v is not None for v in objective_coefficients)
        and len(resources) >= 1
        and all(r["rhs"] is not None for r in resources)
        and all(all(v is not None for v in r["coefficients"]) for r in resources)
    )

    constraints = []
    if ready:
        constraints = [
            {"coefficients": r["coefficients"], "operator": r["operator"], "rhs": r["rhs"]}
            for r in resources
        ]

    model_objective = None
    if len(objective_coefficients) == 2 and all(v is not None for v in objective_coefficients):
        model_objective = {
            "type": "maximize" if optimization == "max" else "minimize",
            "coefficients": objective_coefficients,
            "x1": objective_coefficients[0],
            "x2": objective_coefficients[1],
            "expression": f"Z = {objective_coefficients[0]}X1 + {objective_coefficients[1]}X2",
        }

    return {
        "status": "success" if ready else "needs_confirmation",
        "confidence": "high" if ready else "low",
        "problem_statement": text,
        "optimization": optimization,
        "decision_variables": (
            [
                {"name": "X1", "meaning": f"Number of {products[0]} produced"},
                {"name": "X2", "meaning": f"Number of {products[1]} produced"},
            ] if len(products) == 2 else []
        ),
        "products": products,
        "objective": model_objective,
        "constraints": constraints,
        "resources": resources,
        "message": (
            "Mathematical model generated. Review it before solving."
            if ready else
            "Could not parse the statement with high confidence. Please confirm the coefficients manually."
        ),
    }


@app.route("/api/lpp/parse-word-problem", methods=["POST"])
def parse_word_problem():
    try:
        data = request.get_json(silent=True)
        statement = data.get("problem_statement", "") if data else ""
        if not str(statement).strip():
            return jsonify({"status": "error", "message": "Problem statement is required."}), 400
        return jsonify(_parse_word_problem(statement)), 200
    except ValueError as error:
        return jsonify({"status": "error", "message": str(error)}), 400
    except Exception as error:
        return jsonify({"status": "error", "message": f"An unexpected error occurred: {str(error)}"}), 500


# ============================================================
# LPP - EXCEL PRACTICAL EXPORT
# ============================================================

@app.route("/api/lpp/export-excel", methods=["POST"])
def export_lpp_excel():
    """Create one professional Excel practical sheet with tables and embedded graph."""
    import os
    import tempfile
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    try:
        data = request.get_json(silent=True) or {}

        raw_objective = data.get("objective")
        if isinstance(raw_objective, dict):
            objective = [float(raw_objective.get("x1", 0)), float(raw_objective.get("x2", 0))]
        elif isinstance(raw_objective, (list, tuple)):
            objective = [float(v) for v in raw_objective]
        else:
            return jsonify({"status": "error", "message": "Objective must contain X1 and X2 coefficients."}), 400

        if len(objective) != 2:
            return jsonify({"status": "error", "message": "This Excel practical supports X1 and X2 only."}), 400

        optimization = str(data.get("optimization", "max")).lower()
        if optimization not in ("max", "min"):
            optimization = "max"

        problem_statement = str(data.get("problem_statement") or "LPP entered using structured coefficients.").strip()
        raw_constraints = data.get("constraints")
        if not isinstance(raw_constraints, list) or not raw_constraints:
            return jsonify({"status": "error", "message": "At least one constraint is required."}), 400

        constraints = []
        for i, c in enumerate(raw_constraints, 1):
            if not isinstance(c, dict):
                return jsonify({"status": "error", "message": f"Constraint {i} is invalid."}), 400
            coeffs = c.get("coefficients")
            if coeffs is None:
                coeffs = [c.get("x1", 0), c.get("x2", 0)]
            if len(coeffs) < 2:
                return jsonify({"status": "error", "message": f"Constraint {i} must contain X1 and X2 coefficients."}), 400
            try:
                c1, c2 = float(coeffs[0]), float(coeffs[1])
                rhs = float(c.get("rhs", c.get("value", 0)))
            except (TypeError, ValueError):
                return jsonify({"status": "error", "message": f"Constraint {i} contains a non-numeric value."}), 400
            op = str(c.get("operator", c.get("relation", "<="))).strip()
            if op not in ("<=", ">=", "="):
                return jsonify({"status": "error", "message": f"Constraint {i} has an invalid operator."}), 400
            constraints.append({"coefficients": [c1, c2], "operator": op, "rhs": rhs})

        result = solve_graphical_lpp(
            objective_coefficients=objective,
            constraints=constraints,
            optimization=optimization,
        )
        if not isinstance(result, dict):
            return jsonify({"status": "error", "message": "The LPP solver returned an invalid result."}), 400

        def xy(point):
            if isinstance(point, dict):
                return float(point.get("x", point.get("x1"))), float(point.get("y", point.get("x2")))
            return float(point[0]), float(point[1])

        def clean_points(values):
            out = []
            for point in values or []:
                try:
                    x, y = xy(point)
                    if not any(abs(x-a) < 1e-9 and abs(y-b) < 1e-9 for a, b in out):
                        out.append((x, y))
                except (TypeError, ValueError, IndexError, KeyError):
                    continue
            return out

        corner_xy = clean_points(result.get("corner_points"))
        feasible_xy = clean_points(result.get("feasible_region"))

        sol = result.get("optimal_solution") or {}
        if isinstance(sol, dict):
            ox = sol.get("x1", sol.get("X1", sol.get("x")))
            oy = sol.get("x2", sol.get("X2", sol.get("y")))
        elif isinstance(sol, (list, tuple)):
            ox = sol[0] if len(sol) > 0 else None
            oy = sol[1] if len(sol) > 1 else None
        else:
            ox = oy = None

        if ox is not None:
            ox = float(ox)
        if oy is not None:
            oy = float(oy)

        oz = result.get("optimal_value", result.get("optimal_z"))
        if oz is not None:
            oz = float(oz)
        elif ox is not None and oy is not None:
            oz = objective[0] * ox + objective[1] * oy

        def fmt(value):
            try:
                value = float(value)
                return str(int(value)) if value.is_integer() else f"{value:g}"
            except (TypeError, ValueError):
                return str(value)

        def equation(a, b, op=None, rhs=None):
            sign = "+" if b >= 0 else "-"
            text = f"{fmt(a)}X1 {sign} {fmt(abs(b))}X2"
            return text if op is None else f"{text} {op} {fmt(rhs)}"

        # ------------------------------------------------------------
        # Workbook / professional Excel styling
        # ------------------------------------------------------------
        wb = Workbook()
        ws = wb.active
        ws.title = "OR LPP Practical"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A5"

        navy = "17365D"
        blue = "2F75B5"
        light_blue = "D9EAF7"
        very_light = "F4F8FB"
        green = "E2F0D9"
        dark_text = "1F2937"
        border_color = "B7C9D6"
        white = "FFFFFF"

        title_font = Font(name="Calibri", size=16, bold=True, color=white)
        subtitle_font = Font(name="Calibri", size=11, italic=True, color="44546A")
        section_font = Font(name="Calibri", size=12, bold=True, color=white)
        header_font = Font(name="Calibri", size=11, bold=True, color=dark_text)
        body_font = Font(name="Calibri", size=11, color=dark_text)
        bold_font = Font(name="Calibri", size=11, bold=True, color=dark_text)
        answer_font = Font(name="Calibri", size=12, bold=True, color=dark_text)

        title_fill = PatternFill("solid", fgColor=navy)
        section_fill = PatternFill("solid", fgColor=blue)
        header_fill = PatternFill("solid", fgColor=light_blue)
        alt_fill = PatternFill("solid", fgColor=very_light)
        answer_fill = PatternFill("solid", fgColor=green)
        thin = Side(style="thin", color=border_color)
        medium = Side(style="medium", color=blue)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        section_border = Border(left=medium, right=medium, top=medium, bottom=medium)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

        widths = {
            "A": 16, "B": 18, "C": 18, "D": 18, "E": 16,
            "F": 16, "G": 4, "H": 16, "I": 16, "J": 16,
            "K": 16, "L": 16,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        # Header area
        ws.merge_cells("A1:L1")
        ws["A1"] = "OPERATION RESEARCH SMART SOLVER"
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:L2")
        ws["A2"] = "Linear Programming Problem — Practical Solution"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = center
        ws.row_dimensions[2].height = 22

        row = 4

        def section(title):
            nonlocal row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
            cell = ws.cell(row, 1, title)
            cell.fill = section_fill
            cell.font = section_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = section_border
            ws.row_dimensions[row].height = 24
            row += 1

        def merged_text(value, font=body_font, min_height=30):
            nonlocal row
            value = str(value)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
            cell = ws.cell(row, 1, value)
            cell.font = font
            cell.alignment = left_wrap
            cell.border = border
            # Merged Excel cells do not auto-fit height, so calculate it ourselves.
            lines = max(1, (len(value) // 125) + value.count("\n") + 1)
            ws.row_dimensions[row].height = max(min_height, min(240, 18 * lines))
            row += 1

        def add_table(headers, rows, table_name, start_col=1, header_fill_color=light_blue):
            nonlocal row
            start_row = row
            end_col = start_col + len(headers) - 1
            for j, header in enumerate(headers, start_col):
                cell = ws.cell(row, j, header)
                cell.fill = PatternFill("solid", fgColor=header_fill_color)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
            ws.row_dimensions[row].height = 28
            row += 1
            for ridx, values in enumerate(rows):
                for j, value in enumerate(values, start_col):
                    cell = ws.cell(row, j, value)
                    cell.font = body_font
                    cell.alignment = center
                    cell.border = border
                    if ridx % 2 == 1:
                        cell.fill = alt_fill
                ws.row_dimensions[row].height = 23
                row += 1
            if rows:
                ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{row-1}"
                table = Table(displayName=table_name, ref=ref)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                ws.add_table(table)
            row += 1

        # 1. Problem statement
        section("1. Problem Statement")
        merged_text(problem_statement, body_font, 54)

        # 2. Mathematical formulation
        section("2. Mathematical Formulation")
        merged_text(
            f"{'Maximize' if optimization == 'max' else 'Minimize'}  Z = {equation(objective[0], objective[1])}",
            answer_font,
            30,
        )
        merged_text("Subject To", bold_font, 25)
        for i, c in enumerate(constraints, 1):
            merged_text(f"C{i}:  {equation(c['coefficients'][0], c['coefficients'][1], c['operator'], c['rhs'])}", body_font, 25)
        merged_text("Non-negativity:  X1 >= 0,  X2 >= 0", body_font, 25)

        # 3. Objective coefficient table
        section("3. Objective Function Coefficients")
        add_table(
            ["Variable", "X1", "X2", "Optimization", "Objective Expression"],
            [["Coefficient", objective[0], objective[1], "Maximize" if optimization == "max" else "Minimize", equation(objective[0], objective[1])]],
            "ObjectiveCoefficients",
            start_col=1,
        )

        # 4. Constraint coefficient table
        section("4. Constraint Coefficient Table")
        constraint_rows = []
        for i, c in enumerate(constraints, 1):
            constraint_rows.append([
                f"C{i}",
                c["coefficients"][0],
                c["coefficients"][1],
                c["operator"],
                c["rhs"],
                equation(c["coefficients"][0], c["coefficients"][1], c["operator"], c["rhs"]),
            ])
        add_table(
            ["Constraint", "X1 Coefficient", "X2 Coefficient", "Relation", "RHS", "Complete Constraint"],
            constraint_rows,
            "ConstraintCoefficients",
            start_col=1,
        )

        # 5. Corner point analysis
        section("5. Corner Point Analysis")
        point_names = ["O"] + [chr(ord("A") + i) for i in range(max(0, len(corner_xy) - 1))]
        corner_rows = []
        for i, (x1, x2) in enumerate(corner_xy):
            label = point_names[i] if i < len(point_names) else f"P{i+1}"
            z_value = objective[0] * x1 + objective[1] * x2
            corner_rows.append([label, x1, x2, z_value])
        if corner_rows:
            add_table(["Point", "X1", "X2", "Objective Z"], corner_rows, "CornerPointAnalysis", start_col=1)
        else:
            merged_text("No corner points were returned by the graphical solver.", body_font, 30)

        # 6. Final answer
        section("6. Final Optimal Solution")
        final_rows = [
            ["Optimization", "Maximize" if optimization == "max" else "Minimize"],
            ["X1", fmt(ox) if ox is not None else "—"],
            ["X2", fmt(oy) if oy is not None else "—"],
            ["Optimal Z", fmt(oz) if oz is not None else "—"],
        ]
        final_start = row
        for label, value in final_rows:
            ws.cell(row, 1, label).font = answer_font
            ws.cell(row, 2, value).font = answer_font
            for col in range(1, 3):
                ws.cell(row, col).fill = answer_fill
                ws.cell(row, col).border = border
                ws.cell(row, col).alignment = center
            ws.row_dimensions[row].height = 25
            row += 1
        row += 1

        # 7. Graph — use a real temporary PNG path so openpyxl embeds it reliably.
        section("7. Graphical Solution — Feasible Region")
        graph_row = row
        graph_path = None
        try:
            max_x = max_y = 10.0
            all_points = corner_xy + feasible_xy + ([(ox, oy)] if ox is not None and oy is not None else [])
            for x, y in all_points:
                max_x = max(max_x, abs(x) * 1.25 + 1)
                max_y = max(max_y, abs(y) * 1.25 + 1)
            max_x = min(max_x, 1000)
            max_y = min(max_y, 1000)

            fig, ax = plt.subplots(figsize=(11, 7), dpi=150)

            # Feasible region
            if len(feasible_xy) >= 3:
                poly = list(feasible_xy)
                if poly[0] != poly[-1]:
                    poly.append(poly[0])
                px = [p[0] for p in poly]
                py = [p[1] for p in poly]
                ax.fill(px, py, alpha=0.22, label="Feasible Region")
                ax.plot(px, py, linewidth=1.8)

            # Constraint lines
            for i, c in enumerate(constraints, 1):
                a, b = c["coefficients"]
                rhs = c["rhs"]
                if abs(b) > 1e-12:
                    y0 = rhs / b
                    y1 = (rhs - a * max_x) / b
                    ax.plot([0, max_x], [y0, y1], linewidth=1.5, label=f"C{i}")
                elif abs(a) > 1e-12:
                    ax.axvline(rhs / a, linewidth=1.5, label=f"C{i}")

            # Corner points with O/A/B/C labels
            if corner_xy:
                ax.scatter(
                    [p[0] for p in corner_xy],
                    [p[1] for p in corner_xy],
                    s=65,
                    zorder=5,
                    label="Corner Points",
                )
                for i, (x, y) in enumerate(corner_xy):
                    label = point_names[i] if i < len(point_names) else f"P{i+1}"
                    ax.annotate(
                        label,
                        (x, y),
                        xytext=(7, 7),
                        textcoords="offset points",
                        fontsize=11,
                        fontweight="bold",
                    )

            # Optimal solution
            if ox is not None and oy is not None:
                ax.scatter([ox], [oy], s=180, marker="*", zorder=7, label="Optimal Solution")
                ax.annotate(
                    "Optimal",
                    (ox, oy),
                    xytext=(10, -18),
                    textcoords="offset points",
                    fontsize=10,
                    fontweight="bold",
                )

            ax.set_xlim(0, max_x)
            ax.set_ylim(0, max_y)
            ax.set_xlabel("X1", fontsize=11)
            ax.set_ylabel("X2", fontsize=11)
            ax.set_title("LPP Graphical Solution", fontsize=14, fontweight="bold", pad=12)
            ax.grid(True, alpha=0.25)
            ax.legend(loc="best", fontsize=9)
            fig.tight_layout()

            temp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            graph_path = temp.name
            temp.close()
            fig.savefig(graph_path, format="png", bbox_inches="tight", facecolor="white")
            plt.close(fig)

            image = XLImage(graph_path)
            image.width = 900
            image.height = 570
            ws.add_image(image, f"A{graph_row}")

            # Reserve enough visible worksheet space for the image.
            for rr in range(graph_row, graph_row + 32):
                ws.row_dimensions[rr].height = 18
            row += 32
        except Exception as graph_error:
            plt.close("all")
            merged_text(f"Graph could not be embedded: {graph_error}", body_font, 42)

        # 8. Feasible region coordinates
        section("8. Feasible Region Coordinates")
        feasible_rows = [[f"F{i}", x, y] for i, (x, y) in enumerate(feasible_xy, 1)]
        if feasible_rows:
            add_table(["Point", "X1", "X2"], feasible_rows, "FeasibleCoordinates", start_col=1)
        else:
            merged_text("No feasible-region coordinates were returned.", body_font, 30)

        # Workbook polish
        for r in range(1, row + 1):
            for c in range(1, 13):
                cell = ws.cell(r, c)
                if cell.value is not None and cell.border == Border():
                    cell.border = border

        ws.auto_filter.ref = None
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.print_title_rows = "1:2"
        ws.print_area = f"A1:L{row}"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        if graph_path:
            try:
                os.remove(graph_path)
            except OSError:
                pass

        return send_file(
            output,
            as_attachment=True,
            download_name="OR_LPP_Practical.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except ValueError as error:
        return jsonify({"status": "error", "message": str(error)}), 400
    except Exception as error:
        return jsonify({"status": "error", "message": f"An unexpected error occurred: {str(error)}"}), 500


# ============================================================
# TRANSPORTATION - NORTH WEST CORNER
# ============================================================

@app.route("/api/transportation/north-west", methods=["POST"])
def solve_transportation_north_west():
    try:
        data = request.get_json()

        if not data:
            return jsonify({
                "status": "error",
                "message": "No JSON data was provided."
            }), 400

        costs = data.get("costs")
        supply = data.get("supply")
        demand = data.get("demand")

        if costs is None or supply is None or demand is None:
            return jsonify({
                "status": "error",
                "message": "Costs, supply and demand are required."
            }), 400

        if sum(supply) != sum(demand):
            return jsonify({
                "status": "error",
                "message": "Total supply and total demand must be equal."
            }), 400

        result = north_west_corner(
            costs=costs,
            supply=supply,
            demand=demand
        )

        return jsonify({
            "status": "success",
            "method": "North-West Corner",
            "result": result
        }), 200

    except ValueError as error:
        return jsonify({
            "status": "error",
            "message": str(error)
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 500


# ============================================================
# TRANSPORTATION - LEAST COST METHOD
# ============================================================

@app.route("/api/transportation/least-cost", methods=["POST"])
def solve_transportation_least_cost():
    try:
        data = request.get_json()

        if not data:
            return jsonify({
                "status": "error",
                "message": "No JSON data was provided."
            }), 400

        costs = data.get("costs")
        supply = data.get("supply")
        demand = data.get("demand")

        if costs is None or supply is None or demand is None:
            return jsonify({
                "status": "error",
                "message": "Costs, supply and demand are required."
            }), 400

        if sum(supply) != sum(demand):
            return jsonify({
                "status": "error",
                "message": "Total supply and total demand must be equal."
            }), 400

        result = least_cost_method(
            costs=costs,
            supply=supply,
            demand=demand
        )

        return jsonify({
            "status": "success",
            "method": "Least Cost Method",
            "result": result
        }), 200

    except ValueError as error:
        return jsonify({
            "status": "error",
            "message": str(error)
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 500


# ============================================================
# TRANSPORTATION - VOGEL'S APPROXIMATION METHOD
# ============================================================

@app.route("/api/transportation/vogel", methods=["POST"])
def solve_transportation_vogel():
    try:
        data = request.get_json()

        if not data:
            return jsonify({
                "status": "error",
                "message": "No JSON data was provided."
            }), 400

        costs = data.get("costs")
        supply = data.get("supply")
        demand = data.get("demand")

        if costs is None or supply is None or demand is None:
            return jsonify({
                "status": "error",
                "message": "Costs, supply and demand are required."
            }), 400

        if sum(supply) != sum(demand):
            return jsonify({
                "status": "error",
                "message": "Total supply and total demand must be equal."
            }), 400

        result = vogel_approximation_method(
            costs=costs,
            supply=supply,
            demand=demand
        )

        return jsonify({
            "status": "success",
            "method": "Vogel's Approximation Method",
            "result": result
        }), 200

    except ValueError as error:
        return jsonify({
            "status": "error",
            "message": str(error)
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 500


# ============================================================
# TRANSPORTATION - MODI METHOD
# ============================================================

@app.route("/api/transportation/modi", methods=["POST"])
def solve_transportation_modi():
    try:
        data = request.get_json()

        if not data:
            return jsonify({
                "status": "error",
                "message": "No JSON data was provided."
            }), 400

        costs = data.get("costs")
        supply = data.get("supply")
        demand = data.get("demand")

        if costs is None or supply is None or demand is None:
            return jsonify({
                "status": "error",
                "message": "Costs, supply and demand are required."
            }), 400

        if sum(supply) != sum(demand):
            return jsonify({
                "status": "error",
                "message": "Total supply and total demand must be equal."
            }), 400

        result = modi_method(
            costs=costs,
            supply=supply,
            demand=demand
        )

        return jsonify({
            "status": "success",
            "method": "MODI Method",
            "result": result
        }), 200

    except ValueError as error:
        return jsonify({
            "status": "error",
            "message": str(error)
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 400

    except Exception as error:
        return jsonify({
            "status": "error",
            "message": f"An unexpected error occurred: {str(error)}"
        }), 500


if __name__ == "__main__":
    app.run(debug=True)